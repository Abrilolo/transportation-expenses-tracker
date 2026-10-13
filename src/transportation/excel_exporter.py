from __future__ import annotations

from pathlib import Path

from .suica_record import SuicaHistoryRecord


HEADERS = (
    "Date",
    "Process",
    "Terminal",
    "Entry station",
    "Entry station (Japanese)",
    "Entry line",
    "Exit station",
    "Exit station (Japanese)",
    "Exit line",
    "Expense (JPY)",
    "Charge (JPY)",
    "Balance (JPY)",
    "Sequence",
    "Raw data",
)


def export_to_excel(path: Path, records: list[SuicaHistoryRecord]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("Excel export requires openpyxl: pip install openpyxl") from exc

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Transactions"
    worksheet.freeze_panes = "A2"
    worksheet.append(HEADERS)

    for record in records:
        worksheet.append(
            (
                record.used_at,
                record.process_name,
                record.terminal_name,
                record.entry_station.display_name if record.entry_station else None,
                record.entry_station.station_name if record.entry_station else None,
                record.entry_station.line_name if record.entry_station else None,
                record.exit_station.display_name if record.exit_station else None,
                record.exit_station.station_name if record.exit_station else None,
                record.exit_station.line_name if record.exit_station else None,
                record.expense_yen,
                record.charge_yen,
                record.balance_yen,
                record.sequence_number,
                record.raw_hex,
            )
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    widths = (12, 22, 30, 24, 24, 22, 24, 24, 22, 15, 15, 15, 12, 50)
    for column, width in zip(worksheet.columns, widths):
        worksheet.column_dimensions[column[0].column_letter].width = width

    if records:
        table = Table(displayName="SuicaTransactions", ref=f"A1:N{len(records) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
